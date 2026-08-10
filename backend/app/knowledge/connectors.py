import base64
import ctypes
import hashlib
import os
import re
import shutil
import stat
import subprocess
import zipfile
from collections import deque
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import Callable, Iterator
from urllib.parse import urlsplit

from app.knowledge.credentials import (
    KnowledgeCredentialStore,
    SourceCredential,
)
from app.knowledge.extractors import (
    SUPPORTED_EXTENSIONS,
    extract_text_with_metadata,
    normalize_text,
)
from app.knowledge.ocr import TesseractOcrEngine
from app.knowledge.processing_policy import (
    classify_file,
    path_semantic_text,
)
from app.knowledge.shortcuts import (
    SHORTCUT_PARSER_VERSION,
    ShortcutParseError,
    parse_shortcut,
    shortcut_semantic_text,
)
from app.policies.command_policy import CommandPolicyService


SOURCE_TYPES = {
    "local_directory",
    "network_share",
    "git",
    "gitlab",
    "svn",
}
EXCLUDED_PARTS = {
    ".git",
    ".idea",
    ".pytest_cache",
    ".svn",
    ".venv",
    "build",
    "coverage",
    "dist",
    "node_modules",
    "target",
    "vendor",
}


def _path_for_io(
    path: Path,
    *,
    platform_name: str | None = None,
) -> Path:
    if (platform_name or os.name) != "nt":
        return path
    raw = str(path)
    if raw.startswith("\\\\?\\") or len(raw) < 248:
        return path
    if raw.startswith("\\\\"):
        return Path("\\\\?\\UNC\\" + raw.lstrip("\\"))
    return Path("\\\\?\\" + raw)


@dataclass(frozen=True)
class SourceConfig:
    id: str
    source_type: str
    location: str
    reference: str | None
    subpath: str | None
    credential_ref: str | None


@dataclass(frozen=True)
class CollectedDocument:
    path: str
    text: str
    content_hash: str
    language: str
    encoding: str
    processing_mode: str
    extractor: str
    extractor_version: str | None = None
    processor_variant: str | None = None


@dataclass(frozen=True)
class ReusableFile:
    file_size: int | None
    modified_at: datetime | None
    processing_status: str
    reason_code: str | None
    raw_content_hash: str | None
    has_document: bool


@dataclass(frozen=True)
class CollectionResult:
    revision: str | None
    documents: list[CollectedDocument]
    files_seen: int
    rejected_files: int
    skipped_files: int
    duplicate_files: int
    reused_paths: tuple[str, ...] = ()


@dataclass(frozen=True)
class CollectionRejection:
    relative_path: str
    entry_kind: str
    disposition: str
    extension: str
    file_size: int | None
    reason_code: str
    extractor: str
    extractor_version: str | None = None
    error_type: str | None = None
    error_message: str | None = None


@dataclass(frozen=True)
class CollectionObservation:
    relative_path: str
    entry_kind: str
    extension: str
    file_size: int | None
    modified_at: datetime | None
    processing_mode: str
    processing_status: str
    reason_code: str | None
    raw_content_hash: str | None = None


@dataclass(frozen=True)
class ValidationResult:
    ok: bool
    revision: str | None
    message: str


CollectionProgress = Callable[[dict[str, int | str]], None]
CollectionRejectionSink = Callable[[CollectionRejection], None]
CollectionObservationSink = Callable[[CollectionObservation], None]


def _path_identity(path: Path) -> str:
    absolute = os.path.normcase(os.path.abspath(os.fspath(path)))
    anchor = os.path.normcase(Path(absolute).anchor)
    if anchor and absolute.rstrip("\\/") == anchor.rstrip("\\/"):
        return anchor
    return absolute.rstrip("\\/")


def _contains_path(parent: Path, child: Path) -> bool:
    parent_identity = _path_identity(parent)
    child_identity = _path_identity(child)
    try:
        return (
            os.path.commonpath([parent_identity, child_identity])
            == parent_identity
        )
    except ValueError:
        return False


class SourceConnectorManager:
    def __init__(
        self,
        *,
        cache_root: Path,
        allowed_roots: list[Path],
        credential_store: KnowledgeCredentialStore,
        command_policy: CommandPolicyService,
        git_executable: str,
        svn_executable: str,
        max_file_bytes: int,
        max_spreadsheet_cells: int = 250_000,
        ocr_engine: TesseractOcrEngine | None = None,
    ) -> None:
        self._cache_root = cache_root.resolve()
        self._allowed_roots = allowed_roots
        self._credentials = credential_store
        self._policy = command_policy
        self._git = git_executable
        self._svn = svn_executable
        self._max_file_bytes = max_file_bytes
        self._max_spreadsheet_cells = max_spreadsheet_cells
        self._ocr_engine = ocr_engine

    @staticmethod
    def normalized_source_key(
        *,
        source_type: str,
        location: str,
        reference: str | None,
        subpath: str | None,
        scope: str,
    ) -> str:
        normalized_location = location.strip().rstrip("/\\")
        if source_type in {"local_directory", "network_share"}:
            normalized_location = normalized_location.casefold()
        payload = "\n".join(
            (
                source_type,
                normalized_location,
                (reference or "").strip(),
                (subpath or "").replace("\\", "/").strip("/"),
                scope,
            )
        )
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()

    def validate_definition(
        self,
        *,
        source_type: str,
        location: str,
        subpath: str | None,
    ) -> None:
        if source_type not in SOURCE_TYPES:
            raise ValueError("Unsupported knowledge source type")
        if not location.strip():
            raise ValueError("Knowledge source location is required")
        self._validate_subpath(subpath)
        if source_type == "local_directory":
            self._resolve_allowed_local_root(location)
        elif source_type == "network_share":
            if os.name != "nt" or not location.startswith("\\\\"):
                raise ValueError(
                    "Network share must be a Windows UNC path"
                )
        else:
            self._validate_repository_location(source_type, location)

    def validate(self, source: SourceConfig) -> ValidationResult:
        self.validate_definition(
            source_type=source.source_type,
            location=source.location,
            subpath=source.subpath,
        )
        credential = self._credentials.get(source.credential_ref)
        if source.source_type == "local_directory":
            root = self._selected_root(
                self._resolve_allowed_local_root(source.location),
                source.subpath,
            )
            return ValidationResult(
                True,
                self._git_revision(root),
                f"Directory is readable: {root.name}",
            )
        if source.source_type == "network_share":
            with self._network_connection(source.location, credential):
                root = self._selected_root(Path(source.location), source.subpath)
                if not root.is_dir():
                    raise ValueError("Network share directory is unavailable")
                return ValidationResult(
                    True,
                    None,
                    "Network share is readable",
                )
        if source.source_type in {"git", "gitlab"}:
            revision = self._git_remote_revision(source, credential)
            return ValidationResult(
                True, revision, "Git repository is reachable"
            )
        revision = self._svn_revision(source, credential)
        return ValidationResult(
            True, revision, "SVN repository is reachable"
        )

    def collect(
        self,
        source: SourceConfig,
        progress: CollectionProgress | None = None,
        rejection: CollectionRejectionSink | None = None,
        observation: CollectionObservationSink | None = None,
        reusable_files: dict[str, ReusableFile] | None = None,
    ) -> CollectionResult:
        credential = self._credentials.get(source.credential_ref)
        if source.source_type == "local_directory":
            root = self._selected_root(
                self._resolve_allowed_local_root(source.location),
                source.subpath,
            )
            return self._read_documents(
                root,
                self._git_revision(root),
                progress,
                rejection,
                observation,
                reusable_files,
            )
        if source.source_type == "network_share":
            with self._network_connection(source.location, credential):
                root = self._selected_root(Path(source.location), source.subpath)
                return self._read_documents(
                    root,
                    None,
                    progress,
                    rejection,
                    observation,
                    reusable_files,
                )
        if source.source_type in {"git", "gitlab"}:
            root, revision = self._materialize_git(source, credential)
        else:
            root, revision = self._materialize_svn(source, credential)
        return self._read_documents(
            self._selected_root(root, source.subpath),
            revision,
            progress,
            rejection,
            observation,
            reusable_files,
        )

    def purge(self, source_id: str) -> None:
        source_cache = self._source_cache_path(source_id)
        resolved_cache = self._cache_root.resolve()
        resolved_source = source_cache.resolve()
        if resolved_cache not in resolved_source.parents:
            raise ValueError("Knowledge source cache path is invalid")
        if resolved_source.is_dir():
            shutil.rmtree(
                resolved_source,
                onexc=self._remove_readonly,
            )

    @staticmethod
    def _remove_readonly(function, path: str, error: BaseException) -> None:
        if isinstance(error, PermissionError):
            os.chmod(path, stat.S_IWRITE)
            function(path)
            return
        raise error

    def _read_documents(
        self,
        root: Path,
        revision: str | None,
        progress: CollectionProgress | None = None,
        rejection: CollectionRejectionSink | None = None,
        observation: CollectionObservationSink | None = None,
        reusable_files: dict[str, ReusableFile] | None = None,
    ) -> CollectionResult:
        if not root.is_dir():
            raise ValueError("Selected source subpath does not exist")
        documents: list[CollectedDocument] = []
        rejected = 0
        skipped = 0
        duplicates = 0
        files_discovered = 0
        files_processed = 0
        directories_scanned = 0
        pending_directories: deque[tuple[Path, PurePosixPath]] = deque(
            [(root, PurePosixPath())]
        )
        visited_directories: set[str] = set()
        coverage_roots = [root]
        allowed_shortcut_root = root
        reusable_files = reusable_files or {}
        reused_paths: list[str] = []

        def append_path_document(
            path: Path,
            reason_code: str,
            *,
            relative_path: str | None = None,
            text: str | None = None,
            extractor: str = "path-semantic",
            extractor_version: str | None = None,
        ) -> None:
            relative_path = relative_path or path.relative_to(root).as_posix()
            path_text = path_semantic_text(
                relative_path,
                reason_code=reason_code,
            ) if text is None else text
            documents.append(
                CollectedDocument(
                    path=relative_path,
                    text=path_text,
                    content_hash=hashlib.sha256(
                        path_text.encode("utf-8")
                    ).hexdigest(),
                    language="path",
                    encoding="path-metadata",
                    processing_mode="path_only",
                    extractor=extractor,
                    extractor_version=extractor_version,
                )
            )

        if root.anchor.startswith("\\\\"):
            allowed_shortcut_root = Path(root.anchor)
        else:
            for candidate in self._allowed_roots:
                resolved = candidate.resolve()
                if _contains_path(resolved, root):
                    allowed_shortcut_root = resolved
                    break

        def schedule_shortcut_directory(
            target: Path,
            logical_path: PurePosixPath,
        ) -> bool:
            nonlocal coverage_roots
            if any(
                _contains_path(value, target) or _contains_path(target, value)
                for value in coverage_roots
            ):
                return False
            coverage_roots.append(target)
            pending_directories.append((target, logical_path))
            return True

        def report(
            phase: str,
            directory: Path,
            *,
            logical_path: PurePosixPath | None = None,
            current_directory_files: int = 0,
            error: str = "",
        ) -> None:
            if progress is None:
                return
            relative = (
                logical_path.as_posix()
                if logical_path and logical_path.parts
                else "."
            )
            data: dict[str, int | str] = {
                "phase": phase,
                "directory": relative,
                "directories_scanned": directories_scanned,
                "directories_pending": len(pending_directories),
                "files_discovered": files_discovered,
                "files_processed": files_processed,
                "current_directory_files": current_directory_files,
                "rejected_files": rejected,
                "skipped_files": skipped,
            }
            if error:
                data["error"] = error[:500]
            progress(data)

        while pending_directories:
            directory, logical_directory = pending_directories.popleft()
            identity = _path_identity(directory)
            if identity in visited_directories:
                report(
                    "duplicate_directory_skipped",
                    directory,
                    logical_path=logical_directory,
                )
                continue
            visited_directories.add(identity)
            report("started", directory, logical_path=logical_directory)
            try:
                with os.scandir(directory) as iterator:
                    entries = sorted(
                        iterator,
                        key=lambda entry: entry.name.casefold(),
                    )
            except OSError as exc:
                rejected += 1
                self._report_rejection(
                    rejection,
                    root=root,
                    path=directory,
                    relative_path=(
                        logical_directory.as_posix()
                        if logical_directory.parts
                        else "."
                    ),
                    entry_kind="directory",
                    disposition="rejected",
                    file_size=None,
                    reason_code="directory_read_error",
                    error=exc,
                )
                directories_scanned += 1
                report(
                    "failed",
                    directory,
                    logical_path=logical_directory,
                    error=str(exc),
                )
                continue

            directory_files: list[tuple[Path, str, str]] = []
            for entry in entries:
                path = Path(entry.path)
                relative_path = (
                    logical_directory / entry.name
                ).as_posix()
                try:
                    if entry.is_dir(follow_symlinks=False):
                        directory_stat = entry.stat(follow_symlinks=False)
                        self._report_observation(
                            observation,
                            root=root,
                            path=path,
                            relative_path=relative_path,
                            entry_kind="directory",
                            file_size=None,
                            modified_at=datetime.fromtimestamp(
                                directory_stat.st_mtime,
                                tz=timezone.utc,
                            ),
                            processing_mode="metadata_only",
                            processing_status="observed",
                            reason_code="directory_entry",
                        )
                        if entry.name not in EXCLUDED_PARTS:
                            pending_directories.append(
                                (
                                    Path(entry.path),
                                    logical_directory / entry.name,
                                )
                            )
                        continue
                    if not entry.is_file(follow_symlinks=False):
                        continue
                    file_stat = entry.stat(follow_symlinks=False)
                    file_size = file_stat.st_size
                    modified_at = datetime.fromtimestamp(
                        file_stat.st_mtime,
                        tz=timezone.utc,
                    )
                except OSError as exc:
                    rejected += 1
                    self._report_observation(
                        observation,
                        root=root,
                        path=path,
                        relative_path=relative_path,
                        entry_kind="file",
                        file_size=None,
                        modified_at=None,
                        processing_mode="metadata_only",
                        processing_status="rejected",
                        reason_code="file_stat_error",
                    )
                    self._report_rejection(
                        rejection,
                        root=root,
                        path=path,
                        relative_path=relative_path,
                        entry_kind="file",
                        disposition="rejected",
                        file_size=None,
                        reason_code="file_stat_error",
                        error=exc,
                    )
                    continue
                files_discovered += 1
                decision = classify_file(
                    relative_path,
                    file_size=file_size,
                    max_file_bytes=self._max_file_bytes,
                )
                if path.suffix.casefold() == ".lnk":
                    target_path: str | None = None
                    target_kind: str | None = None
                    target_status = "shortcut_parse_failed"
                    try:
                        raw_content_hash = self._sha256_file(path)
                    except OSError as exc:
                        rejected += 1
                        files_processed += 1
                        self._report_observation(
                            observation,
                            root=root,
                            path=path,
                            relative_path=relative_path,
                            entry_kind="file",
                            file_size=file_size,
                            modified_at=modified_at,
                            processing_mode="path_only",
                            processing_status="rejected",
                            reason_code="raw_hash_read_error",
                        )
                        self._report_rejection(
                            rejection,
                            root=root,
                            path=path,
                            relative_path=relative_path,
                            entry_kind="file",
                            disposition="rejected",
                            file_size=file_size,
                            reason_code="raw_hash_read_error",
                            error=exc,
                        )
                        append_path_document(
                            path,
                            "raw_hash_read_error",
                            relative_path=relative_path,
                        )
                        continue
                    try:
                        parsed = parse_shortcut(path)
                        target_path = parsed.target_path
                        target = Path(target_path)
                        if not _contains_path(allowed_shortcut_root, target):
                            target_status = "shortcut_outside_allowed_root"
                        else:
                            try:
                                target_stat = target.stat()
                            except PermissionError:
                                target_status = "shortcut_auth_denied"
                            except FileNotFoundError:
                                target_status = "shortcut_target_missing"
                            except OSError:
                                target_status = "shortcut_target_unreachable"
                            else:
                                if stat.S_ISDIR(target_stat.st_mode):
                                    target_kind = "directory"
                                    logical_target = (
                                        logical_directory / path.stem
                                    )
                                    target_status = (
                                        "shortcut_target_already_covered"
                                        if any(
                                            _contains_path(value, target)
                                            for value in coverage_roots
                                        )
                                        else (
                                            "shortcut_target_enqueued"
                                            if schedule_shortcut_directory(
                                                target,
                                                logical_target,
                                            )
                                            else "shortcut_target_already_covered"
                                        )
                                    )
                                elif stat.S_ISREG(target_stat.st_mode):
                                    target_kind = "file"
                                    if any(
                                        _contains_path(value, target)
                                        for value in coverage_roots
                                    ):
                                        target_status = (
                                            "shortcut_target_already_covered"
                                        )
                                        semantic_text = shortcut_semantic_text(
                                            relative_path,
                                            target_path=target_path,
                                            target_status=target_status,
                                            target_kind=target_kind,
                                        )
                                        self._report_observation(
                                            observation,
                                            root=root,
                                            path=path,
                                            relative_path=relative_path,
                                            entry_kind="file",
                                            file_size=file_size,
                                            modified_at=modified_at,
                                            processing_mode="path_only",
                                            processing_status="observed",
                                            reason_code=target_status,
                                            raw_content_hash=raw_content_hash,
                                        )
                                        append_path_document(
                                            path,
                                            target_status,
                                            relative_path=relative_path,
                                            text=semantic_text,
                                            extractor="shell-link",
                                            extractor_version=(
                                                SHORTCUT_PARSER_VERSION
                                            ),
                                        )
                                        files_processed += 1
                                        continue
                                    logical_target = (
                                        logical_directory
                                        / path.stem
                                        / target.name
                                    ).as_posix()
                                    target_decision = classify_file(
                                        logical_target,
                                        file_size=target_stat.st_size,
                                        max_file_bytes=self._max_file_bytes,
                                    )
                                    try:
                                        target_raw_hash = self._sha256_file(target)
                                    except PermissionError:
                                        target_status = "shortcut_auth_denied"
                                    except OSError:
                                        target_status = "shortcut_target_unreachable"
                                    else:
                                        files_discovered += 1
                                        self._report_observation(
                                            observation,
                                            root=root,
                                            path=target,
                                            relative_path=logical_target,
                                            entry_kind="file",
                                            file_size=target_stat.st_size,
                                            modified_at=datetime.fromtimestamp(
                                                target_stat.st_mtime,
                                                tz=timezone.utc,
                                            ),
                                            processing_mode=target_decision.mode,
                                            processing_status="observed",
                                            reason_code="shortcut_target_flattened",
                                            raw_content_hash=target_raw_hash,
                                        )
                                        directory_files.append(
                                            (
                                                target,
                                                target_decision.mode,
                                                logical_target,
                                            )
                                        )
                                        target_status = "shortcut_target_enqueued"
                                else:
                                    target_status = "shortcut_target_unsupported"
                    except ShortcutParseError:
                        pass
                    semantic_text = shortcut_semantic_text(
                        relative_path,
                        target_path=target_path,
                        target_status=target_status,
                        target_kind=target_kind,
                    )
                    self._report_observation(
                        observation,
                        root=root,
                        path=path,
                        relative_path=relative_path,
                        entry_kind="file",
                        file_size=file_size,
                        modified_at=modified_at,
                        processing_mode="path_only",
                        processing_status="observed",
                        reason_code=target_status,
                        raw_content_hash=raw_content_hash,
                    )
                    append_path_document(
                        path,
                        target_status,
                        relative_path=relative_path,
                        text=semantic_text,
                        extractor="shell-link",
                        extractor_version=SHORTCUT_PARSER_VERSION,
                    )
                    files_processed += 1
                    continue
                reusable = reusable_files.get(relative_path)
                reusable_metadata_matches = (
                    reusable is not None
                    and reusable.file_size == file_size
                    and reusable.modified_at == modified_at
                    and reusable.raw_content_hash is not None
                )
                if reusable_metadata_matches and (
                    reusable.has_document
                    or reusable.processing_status == "rejected"
                ):
                    provenance_reason = (
                        "shortcut_target_flattened"
                        if not _contains_path(root, path)
                        else reusable.reason_code
                    )
                    self._report_observation(
                        observation,
                        root=root,
                        path=path,
                        relative_path=relative_path,
                        entry_kind="file",
                        file_size=file_size,
                        modified_at=modified_at,
                        processing_mode=decision.mode,
                        processing_status=reusable.processing_status,
                        reason_code=provenance_reason,
                        raw_content_hash=reusable.raw_content_hash,
                    )
                    reused_paths.append(relative_path)
                    files_processed += 1
                    if reusable.processing_status in {
                        "metadata_only",
                        "rejected",
                    }:
                        disposition = (
                            "rejected"
                            if reusable.processing_status == "rejected"
                            else "skipped"
                        )
                        if disposition == "rejected":
                            rejected += 1
                        else:
                            skipped += 1
                        self._report_rejection(
                            rejection,
                            root=root,
                            path=path,
                            relative_path=relative_path,
                            entry_kind="file",
                            disposition=disposition,
                            file_size=file_size,
                            reason_code=(
                                reusable.reason_code or "reused_path_evidence"
                            ),
                        )
                    continue
                try:
                    raw_content_hash = (
                        reusable.raw_content_hash
                        if reusable_metadata_matches
                        else self._sha256_file(path)
                    )
                except OSError as exc:
                    rejected += 1
                    files_processed += 1
                    self._report_observation(
                        observation,
                        root=root,
                        path=path,
                        relative_path=relative_path,
                        entry_kind="file",
                        file_size=file_size,
                        modified_at=modified_at,
                        processing_mode=decision.mode,
                        processing_status="rejected",
                        reason_code="raw_hash_read_error",
                    )
                    self._report_rejection(
                        rejection,
                        root=root,
                        path=path,
                        relative_path=relative_path,
                        entry_kind="file",
                        disposition="rejected",
                        file_size=file_size,
                        reason_code="raw_hash_read_error",
                        error=exc,
                    )
                    append_path_document(
                        path,
                        "raw_hash_read_error",
                        relative_path=relative_path,
                    )
                    continue
                if decision.mode == "metadata_only":
                    self._report_observation(
                        observation,
                        root=root,
                        path=path,
                        relative_path=relative_path,
                        entry_kind="file",
                        file_size=file_size,
                        modified_at=modified_at,
                        processing_mode=decision.mode,
                        processing_status="metadata_only",
                        reason_code=decision.reason_code,
                        raw_content_hash=raw_content_hash,
                    )
                    skipped += 1
                    files_processed += 1
                    self._report_rejection(
                        rejection,
                        root=root,
                        path=path,
                        relative_path=relative_path,
                        entry_kind="file",
                        disposition="skipped",
                        file_size=file_size,
                        reason_code=(
                            decision.reason_code or "metadata_only_policy"
                        ),
                    )
                    append_path_document(
                        path,
                        decision.reason_code or "metadata_only_policy",
                        relative_path=relative_path,
                    )
                    continue
                self._report_observation(
                    observation,
                    root=root,
                    path=path,
                    relative_path=relative_path,
                    entry_kind="file",
                    file_size=file_size,
                    modified_at=modified_at,
                    processing_mode=decision.mode,
                    processing_status=(
                        "metadata_only"
                        if decision.mode == "metadata_only"
                        else "observed"
                    ),
                    reason_code=decision.reason_code,
                    raw_content_hash=raw_content_hash,
                )
                if decision.mode == "path_only":
                    files_processed += 1
                    text = path_semantic_text(
                        relative_path,
                        reason_code=(
                            decision.reason_code
                            or "path_only"
                        ),
                    )
                    documents.append(
                        CollectedDocument(
                            path=relative_path,
                            text=text,
                            content_hash=hashlib.sha256(
                                text.encode("utf-8")
                            ).hexdigest(),
                            language="path",
                            encoding="path-metadata",
                            processing_mode="path_only",
                            extractor="path-semantic",
                        )
                    )
                    continue
                directory_files.append((path, decision.mode, relative_path))

            for path, processing_mode, relative_path in directory_files:
                file_size: int | None = None
                try:
                    io_path = _path_for_io(path)
                    file_size = io_path.stat().st_size
                    extracted = extract_text_with_metadata(
                        io_path,
                        max_spreadsheet_cells=(
                            self._max_spreadsheet_cells
                        ),
                        max_output_characters=self._max_file_bytes,
                        ocr_engine=self._ocr_engine,
                    )
                    text = normalize_text(extracted.text)
                except (
                    UnicodeDecodeError,
                    OSError,
                    RuntimeError,
                    ValueError,
                    zipfile.BadZipFile,
                ) as exc:
                    rejected += 1
                    self._report_rejection(
                        rejection,
                        root=root,
                        path=path,
                        entry_kind="file",
                        disposition="rejected",
                        file_size=file_size,
                        reason_code=self._rejection_reason(path, exc),
                        error=exc,
                    )
                    append_path_document(
                        path,
                        self._rejection_reason(path, exc),
                        relative_path=relative_path,
                    )
                    files_processed += 1
                    continue
                files_processed += 1
                if not text:
                    rejected += 1
                    self._report_rejection(
                        rejection,
                        root=root,
                        path=path,
                        relative_path=relative_path,
                        entry_kind="file",
                        disposition="rejected",
                        file_size=file_size,
                        reason_code="empty_text",
                    )
                    append_path_document(
                        path,
                        "empty_text",
                        relative_path=relative_path,
                    )
                    continue
                content_hash = hashlib.sha256(
                    text.encode("utf-8")
                ).hexdigest()
                documents.append(
                    CollectedDocument(
                        path=relative_path,
                        text=text,
                        content_hash=content_hash,
                        language=path.suffix.lstrip(".").lower() or "text",
                        encoding=extracted.encoding,
                        processing_mode=processing_mode,
                        extractor=extracted.extractor,
                        extractor_version=extracted.extractor_version,
                        processor_variant=extracted.processor_variant,
                    )
                )
            directories_scanned += 1
            report(
                "completed",
                directory,
                logical_path=logical_directory,
                current_directory_files=len(directory_files),
            )

        return CollectionResult(
            revision=revision,
            documents=documents,
            files_seen=files_discovered,
            rejected_files=rejected,
            skipped_files=skipped,
            duplicate_files=duplicates,
            reused_paths=tuple(reused_paths),
        )

    @staticmethod
    def _sha256_file(path: Path) -> str:
        digest = hashlib.sha256()
        with _path_for_io(path).open("rb") as stream:
            while block := stream.read(1024 * 1024):
                digest.update(block)
        return digest.hexdigest()

    @staticmethod
    def _rejection_reason(path: Path, error: Exception) -> str:
        configured_reason = getattr(error, "reason_code", None)
        if isinstance(configured_reason, str):
            return configured_reason
        if isinstance(error, UnicodeDecodeError):
            return "encoding_unsupported"
        if isinstance(error, zipfile.BadZipFile):
            return "office_archive_invalid"
        if isinstance(error, PermissionError):
            return "file_permission_denied"
        if isinstance(error, OSError):
            return "file_read_error"
        if path.suffix.lower() == ".pdf":
            return "pdf_unreadable"
        if isinstance(error, RuntimeError):
            return "extractor_unavailable"
        return "extractor_rejected"

    @staticmethod
    def _extractor_name(path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return "pdf"
        if suffix == ".xlsx":
            return "openpyxl"
        if suffix in {".docx", ".pptx", ".odt"}:
            return "office-xml"
        if suffix in SUPPORTED_EXTENSIONS:
            return "text"
        return "filesystem"

    @classmethod
    def _report_rejection(
        cls,
        sink: CollectionRejectionSink | None,
        *,
        root: Path,
        path: Path,
        relative_path: str | None = None,
        entry_kind: str,
        disposition: str,
        file_size: int | None,
        reason_code: str,
        error: Exception | None = None,
    ) -> None:
        if sink is None:
            return
        relative_path = relative_path or (
            "." if path == root else path.relative_to(root).as_posix()
        )
        error_message = None
        if error is not None:
            error_message = (
                str(error)
                .replace(str(root), "<source-root>")
                .replace("\r", " ")
                .replace("\n", " ")
            )[:1_000]
        sink(
            CollectionRejection(
                relative_path=relative_path,
                entry_kind=entry_kind,
                disposition=disposition,
                extension=path.suffix.lower(),
                file_size=file_size,
                reason_code=reason_code,
                extractor=(
                    "filesystem"
                    if disposition == "skipped"
                    else cls._extractor_name(path)
                ),
                extractor_version=(
                    cls._extractor_version(path)
                    if disposition == "rejected"
                    else None
                ),
                error_type=type(error).__name__ if error is not None else None,
                error_message=error_message,
            )
        )

    @staticmethod
    def _extractor_version(path: Path) -> str | None:
        if path.suffix.lower() != ".xlsx":
            return None
        try:
            import openpyxl
        except ImportError:
            return None
        return openpyxl.__version__

    @staticmethod
    def _report_observation(
        sink: CollectionObservationSink | None,
        *,
        root: Path,
        path: Path,
        relative_path: str | None = None,
        entry_kind: str,
        file_size: int | None,
        modified_at: datetime | None,
        processing_mode: str,
        processing_status: str,
        reason_code: str | None,
        raw_content_hash: str | None = None,
    ) -> None:
        if sink is None:
            return
        relative_path = relative_path or (
            "." if path == root else path.relative_to(root).as_posix()
        )
        sink(
            CollectionObservation(
                relative_path=relative_path,
                entry_kind=entry_kind,
                extension=path.suffix.lower(),
                file_size=file_size,
                modified_at=modified_at,
                processing_mode=processing_mode,
                processing_status=processing_status,
                reason_code=reason_code,
                raw_content_hash=raw_content_hash,
            )
        )

    def _materialize_git(
        self,
        source: SourceConfig,
        credential: SourceCredential | None,
    ) -> tuple[Path, str]:
        revision = self._git_remote_revision(source, credential)
        destination = self._snapshot_path(source.id, revision)
        if not destination.is_dir():
            destination.parent.mkdir(parents=True, exist_ok=True)
            args = [self._git, "clone", "--depth", "1"]
            if source.reference:
                args.extend(["--branch", source.reference])
            args.extend(["--", source.location, str(destination)])
            self._run(args, credential=credential)
        actual = self._run(
            [self._git, "-C", str(destination), "rev-parse", "HEAD"]
        ).stdout.strip()
        return destination, actual

    def _git_remote_revision(
        self,
        source: SourceConfig,
        credential: SourceCredential | None,
    ) -> str:
        ref = source.reference or "HEAD"
        result = self._run(
            [self._git, "ls-remote", "--", source.location, ref],
            credential=credential,
        )
        line = next(
            (item for item in result.stdout.splitlines() if item.strip()),
            "",
        )
        if not line:
            raise ValueError(f"Git reference was not found: {ref}")
        return line.split()[0]

    def _materialize_svn(
        self,
        source: SourceConfig,
        credential: SourceCredential | None,
    ) -> tuple[Path, str]:
        revision = self._svn_revision(source, credential)
        destination = self._snapshot_path(source.id, revision)
        if not destination.is_dir():
            destination.parent.mkdir(parents=True, exist_ok=True)
            args = [
                self._svn,
                "export",
                "--force",
                "--ignore-externals",
                "--non-interactive",
                "--no-auth-cache",
                "-r",
                revision,
            ]
            input_text = self._append_svn_credentials(args, credential)
            args.extend([source.location, str(destination)])
            self._run(args, input_text=input_text)
        return destination, revision

    def _svn_revision(
        self,
        source: SourceConfig,
        credential: SourceCredential | None,
    ) -> str:
        args = [
            self._svn,
            "info",
            "--show-item",
            "revision",
            "--non-interactive",
            "--no-auth-cache",
        ]
        if source.reference:
            args.extend(["-r", source.reference])
        input_text = self._append_svn_credentials(args, credential)
        args.append(source.location)
        return self._run(args, input_text=input_text).stdout.strip()

    @staticmethod
    def _append_svn_credentials(
        args: list[str],
        credential: SourceCredential | None,
    ) -> str | None:
        if credential is None:
            return None
        args.extend(
            [
                "--username",
                credential.username,
                "--password-from-stdin",
            ]
        )
        return credential.secret + "\n"

    def _run(
        self,
        args: list[str],
        *,
        credential: SourceCredential | None = None,
        input_text: str | None = None,
    ) -> subprocess.CompletedProcess[str]:
        executable_name = Path(args[0]).stem
        actions = {"clone", "ls-remote", "rev-parse", "info", "export"}
        action = next((item for item in args[1:] if item in actions), "")
        policy_subject = f"{executable_name} {action}"
        decision = self._policy.evaluate(
            policy_subject, "knowledge_collection"
        )
        if decision.decision != "allow":
            raise PermissionError(decision.reason)
        environment = os.environ.copy()
        environment["GIT_TERMINAL_PROMPT"] = "0"
        if credential is not None and executable_name.lower().startswith("git"):
            token = base64.b64encode(
                f"{credential.username}:{credential.secret}".encode("utf-8")
            ).decode("ascii")
            environment["GIT_CONFIG_COUNT"] = "1"
            environment["GIT_CONFIG_KEY_0"] = "http.extraHeader"
            environment["GIT_CONFIG_VALUE_0"] = (
                f"Authorization: Basic {token}"
            )
        try:
            return subprocess.run(
                args,
                check=True,
                capture_output=True,
                text=True,
                input=input_text,
                env=environment,
                timeout=300,
            )
        except FileNotFoundError as exc:
            raise RuntimeError(
                f"Required connector executable is unavailable: {args[0]}"
            ) from exc
        except subprocess.CalledProcessError as exc:
            message = (
                exc.stderr or exc.stdout or "Connector command failed"
            ).strip()
            raise RuntimeError(message[:1000]) from exc

    def _resolve_allowed_local_root(self, value: str) -> Path:
        resolved = Path(value).resolve()
        if not resolved.is_dir():
            raise ValueError("Knowledge source path does not exist")
        if not any(
            resolved == root or root in resolved.parents
            for root in self._allowed_roots
        ):
            raise ValueError("Knowledge source path is outside configured roots")
        return resolved

    @staticmethod
    def _selected_root(root: Path, subpath: str | None) -> Path:
        if not subpath:
            return root
        selected = (root / Path(*PurePosixPath(subpath).parts)).resolve()
        resolved_root = root.resolve()
        if selected != resolved_root and resolved_root not in selected.parents:
            raise ValueError("Knowledge source subpath escapes the source root")
        return selected

    @staticmethod
    def _validate_subpath(subpath: str | None) -> None:
        if not subpath:
            return
        value = PurePosixPath(subpath.replace("\\", "/"))
        if value.is_absolute() or ".." in value.parts:
            raise ValueError(
                "Knowledge source subpath must stay inside the source"
            )

    def _validate_repository_location(
        self,
        source_type: str,
        location: str,
    ) -> None:
        parsed = urlsplit(location)
        if parsed.username is not None or parsed.password is not None:
            raise ValueError(
                "Credentials must use the operating system credential store"
            )
        if parsed.scheme in {
            "http",
            "https",
            "ssh",
            "svn",
            "svn+ssh",
            "file",
        }:
            return
        if re.match(r"^[^@\s]+@[^:\s]+:.+$", location):
            return
        if source_type in {"git", "gitlab"}:
            self._resolve_allowed_local_root(location)
            return
        raise ValueError("SVN source must use a supported repository URL")

    def _git_revision(self, root: Path) -> str | None:
        try:
            return self._run(
                [self._git, "-C", str(root), "rev-parse", "HEAD"]
            ).stdout.strip()
        except (PermissionError, RuntimeError):
            return None

    def _snapshot_path(self, source_id: str, revision: str) -> Path:
        snapshot_key = hashlib.sha256(
            revision.encode("utf-8")
        ).hexdigest()[:16]
        return self._source_cache_path(source_id) / snapshot_key

    def _source_cache_path(self, source_id: str) -> Path:
        source_key = hashlib.sha256(
            source_id.encode("utf-8")
        ).hexdigest()[:12]
        return self._cache_root / source_key

    @contextmanager
    def _network_connection(
        self,
        location: str,
        credential: SourceCredential | None,
    ) -> Iterator[None]:
        if credential is None:
            yield
            return
        if os.name != "nt":
            raise RuntimeError(
                "Authenticated network shares require a Windows Gateway host"
            )
        remote = self._unc_share_root(location)
        resource = _NETRESOURCE()
        resource.dwType = 1
        resource.lpRemoteName = remote
        result = ctypes.windll.mpr.WNetAddConnection2W(
            ctypes.byref(resource),
            credential.secret,
            credential.username,
            0,
        )
        if result not in {0, 85}:
            raise OSError(result, "Network share authentication failed")
        try:
            yield
        finally:
            ctypes.windll.mpr.WNetCancelConnection2W(remote, 0, False)

    @staticmethod
    def _unc_share_root(location: str) -> str:
        parts = [part for part in location.split("\\") if part]
        if len(parts) < 2:
            raise ValueError("Network share must include server and share name")
        return f"\\\\{parts[0]}\\{parts[1]}"


class _NETRESOURCE(ctypes.Structure):
    _fields_ = [
        ("dwScope", ctypes.c_ulong),
        ("dwType", ctypes.c_ulong),
        ("dwDisplayType", ctypes.c_ulong),
        ("dwUsage", ctypes.c_ulong),
        ("lpLocalName", ctypes.c_wchar_p),
        ("lpRemoteName", ctypes.c_wchar_p),
        ("lpComment", ctypes.c_wchar_p),
        ("lpProvider", ctypes.c_wchar_p),
    ]
